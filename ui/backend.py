"""
backend.py - EVERY call from the UI to the "brain" of Minerva lives in this file.

The pages never talk to LangGraph, the vector DB, or the LLM directly. They only
call the functions below. That means:

  * BACKEND_MODE=mock (default): the functions return realistic FAKE data
    so the whole UI can be demoed without any servers.
  * BACKEND_MODE=openai: the functions call an OpenAI-compatible model directly
    (OpenAI now; SGLang / llama.cpp on JetStream2 later via OPENAI_BASE_URL).
  * Later, teammates replace the bodies (or just flip USE_MOCK_BACKEND to False
    and implement the matching HTTP endpoints) and the UI keeps working unchanged.

Function contract (keep these signatures stable so the UI doesn't break):

    chat_with_tutor(user_id, course, messages)  -> (sources, token_stream)
    plan_resources(user_id, tasks, message)     -> token_stream
    get_feedback(user_id, file, prompt)         -> dict
    upload_course_material(user_id, file, course=None) -> dict
    list_course_materials(user_id)              -> list[dict]
    delete_course_material(user_id, file_id)    -> bool

A "token_stream" is just a Python generator that yields small pieces of text.
The UI feeds it to st.write_stream(), which prints it like ChatGPT typing.

All functions raise BackendError with a FRIENDLY message if something goes wrong;
the pages catch it and show the message instead of crashing.
"""
import itertools
import threading
import time
import uuid
from datetime import datetime
from typing import Iterator

import config


class BackendError(Exception):
    """Raised for any backend problem. The message is safe to show to students."""


# =============================================================================
# MOCK DATA STORE
# =============================================================================
# This dict plays the role of a real database/vector store. It lives in the
# Python SERVER process (shared by all sessions) but is keyed by user_id, so each
# user only ever sees their own files - exactly how a real backend would behave.
# A lock keeps it safe when several users click at the same moment.
_MOCK_MATERIALS: dict[str, list[dict]] = {}
_MOCK_LOCK = threading.Lock()

# Fake "course material" snippets the mock tutor pretends to retrieve.
_DEFAULT_SOURCES = [
    {"file": "Lecture 05 - Containers and Kubernetes.pdf", "page": 12,
     "snippet": "A pod is the smallest deployable unit in Kubernetes..."},
    {"file": "Syllabus.pdf", "page": 3,
     "snippet": "Project proposals are due Oct 5; final reports are due Nov 30."},
    {"file": "Lecture 03 - Cloud Service Models.pptx", "page": 8,
     "snippet": "IaaS, PaaS and SaaS differ in how much the provider manages for you."},
]


def _stream_text(text: str) -> Iterator[str]:
    """Yield `text` a few words at a time with a small delay (fake LLM streaming)."""
    for word in text.split(" "):
        time.sleep(config.MOCK_STREAM_DELAY)
        yield word + " "


def _check_user(user_id: str) -> None:
    """Basic validation shared by all functions."""
    if not user_id:
        raise BackendError("You need to be logged in to do that.")


# =============================================================================
# OPENAI HELPERS (used only when BACKEND_MODE=openai)
# =============================================================================
# This mode talks to any OpenAI-compatible chat API. Today that's OpenAI itself;
# later the same code reaches SGLang / llama.cpp on JetStream2 by changing
# OPENAI_BASE_URL. There's no vector DB yet, so "RAG" here is a stand-in: the
# text of the student's uploaded materials is pasted into the prompt.

_MAX_CHARS_PER_FILE = 20_000   # text kept per uploaded file
_MAX_CONTEXT_CHARS = 40_000    # total course-material text sent per question
_openai_client = None

TUTOR_SYSTEM_PROMPT = """You are Minerva, a tutor for graduate students in {course}.
Guide the student toward understanding instead of handing over answers:
- Ask what they have tried, give hints, and break problems into steps.
- If the question looks like a graded assignment, do NOT give the full solution;
  explain the concepts and offer a next step instead.
- Explain concepts fully when the student asks for an explanation of a concept.
- When course materials are provided below, ground your answer in them and name
  the file you used. If they don't cover the question, say so.
Use Markdown. Keep answers focused."""

PLANNER_SYSTEM_PROMPT = """You are Minerva's resource-management assistant for a graduate student.
Help them pace their work, prioritize deadlines, and plan use of resources such as
study time, professor office hours, and compute time. Today is {today}.
Be concrete: propose dated time blocks, flag conflicts, and keep it realistic.
Use Markdown."""

FEEDBACK_SYSTEM_PROMPT = """You are Minerva, giving a graduate student feedback on a draft.
Be specific, honest, and constructive; refer to concrete parts of the submission.
If a rubric or prompt is provided, evaluate against it.
Reply with ONLY a JSON object with these keys:
  "summary": a 2-3 sentence overall assessment (string),
  "strengths": list of strings,
  "improvements": list of strings,
  "suggestions": list of concrete next steps (strings)."""


def _client():
    """Create the OpenAI client once per server process and reuse it."""
    global _openai_client
    if _openai_client is None:
        if not config.OPENAI_API_KEY:
            raise BackendError("The AI model isn't configured (OPENAI_API_KEY is missing).")
        from openai import OpenAI  # imported here so other modes don't need it

        _openai_client = OpenAI(
            api_key=config.OPENAI_API_KEY,
            base_url=config.OPENAI_BASE_URL,
            timeout=config.REQUEST_TIMEOUT_SECONDS,
        )
    return _openai_client


def _openai_error(err: Exception) -> BackendError:
    """Turn an OpenAI SDK exception into a message that's safe to show students."""
    import openai

    if isinstance(err, openai.AuthenticationError):
        return BackendError("The AI model rejected our API key. Check OPENAI_API_KEY.")
    if isinstance(err, openai.RateLimitError):
        return BackendError("The AI model is busy or out of quota. Please try again shortly.")
    if isinstance(err, (openai.APIConnectionError, openai.APITimeoutError)):
        return BackendError("Can't reach the AI model right now. Please try again shortly.")
    if isinstance(err, openai.APIStatusError):
        return BackendError(f"The AI model returned an error ({err.status_code}).")
    return BackendError("The AI model failed unexpectedly. Please try again.")


def _openai_stream(messages: list[dict]) -> Iterator[str]:
    """
    Start a streaming chat completion and return a generator of text pieces.

    The request is sent BEFORE returning, so connection/auth errors raise
    BackendError here, where the pages already catch it.
    """
    import openai

    try:
        stream = _client().chat.completions.create(
            model=config.OPENAI_MODEL, messages=messages, stream=True,
        )
    except openai.OpenAIError as err:
        raise _openai_error(err) from err

    def _gen() -> Iterator[str]:
        try:
            for chunk in stream:
                if chunk.choices and chunk.choices[0].delta.content:
                    yield chunk.choices[0].delta.content
        except openai.OpenAIError:
            yield "\n\n_(The answer was cut off by a connection problem. Please try again.)_"

    return _gen()


def _extract_text(file) -> str:
    """Best-effort plain text from an uploaded file (txt/md/code, pdf, docx, pptx, ipynb)."""
    import io
    import json

    name = file.name.lower()
    data = file.getvalue()
    try:
        if name.endswith(".pdf"):
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join(
                f"[page {i}]\n{page.extract_text() or ''}"
                for i, page in enumerate(reader.pages, start=1)
            )
        elif name.endswith(".docx"):
            from docx import Document

            text = "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)
        elif name.endswith(".pptx"):
            from pptx import Presentation

            slides = []
            for i, slide in enumerate(Presentation(io.BytesIO(data)).slides, start=1):
                parts = [s.text_frame.text for s in slide.shapes if s.has_text_frame]
                slides.append(f"[slide {i}]\n" + "\n".join(parts))
            text = "\n".join(slides)
        elif name.endswith(".ipynb"):
            cells = json.loads(data).get("cells", [])
            text = "\n\n".join("".join(c.get("source", [])) for c in cells)
        else:
            text = data.decode("utf-8", errors="replace")
    except Exception:
        raise BackendError(f"Couldn't read the text in {file.name}. Is the file valid?")
    return text.strip()[:_MAX_CHARS_PER_FILE]


def _material_context(user_id: str, course: str) -> tuple[str, list[dict]]:
    """Gather this user's uploaded text for a course, capped in size, plus a source list."""
    with _MOCK_LOCK:
        mine = [m for m in _MOCK_MATERIALS.get(user_id, [])
                if m.get("course") in (course, "Unassigned") and m.get("text")]
    blocks, sources, used = [], [], 0
    for m in mine:
        piece = m["text"][: _MAX_CONTEXT_CHARS - used]
        if not piece:
            break
        blocks.append(f"### File: {m['name']}\n{piece}")
        sources.append({"file": m["name"], "page": "—",
                        "snippet": "Full text included as course context."})
        used += len(piece)
    return "\n\n".join(blocks), sources


# =============================================================================
# REAL BACKEND HELPERS (used only when BACKEND_MODE=api)
# =============================================================================
def _http_post(path: str, **kwargs):
    """POST to the real backend and turn network problems into BackendError."""
    import requests  # imported here so mock mode doesn't need it

    try:
        resp = requests.post(
            f"{config.BACKEND_API_URL}{path}",
            timeout=config.REQUEST_TIMEOUT_SECONDS,
            **kwargs,
        )
        resp.raise_for_status()
        return resp
    except requests.exceptions.ConnectionError:
        raise BackendError("Can't reach the Minerva server right now. Please try again shortly.")
    except requests.exceptions.Timeout:
        raise BackendError("The Minerva server took too long to respond. Please try again.")
    except requests.exceptions.HTTPError as err:
        raise BackendError(f"The server returned an error ({err.response.status_code}).")


def _http_get(path: str, **kwargs):
    import requests

    try:
        resp = requests.get(
            f"{config.BACKEND_API_URL}{path}",
            timeout=config.REQUEST_TIMEOUT_SECONDS,
            **kwargs,
        )
        resp.raise_for_status()
        return resp
    except requests.exceptions.RequestException:
        raise BackendError("Can't reach the Minerva server right now. Please try again shortly.")


def _http_delete(path: str, **kwargs):
    import requests

    try:
        resp = requests.delete(
            f"{config.BACKEND_API_URL}{path}",
            timeout=config.REQUEST_TIMEOUT_SECONDS,
            **kwargs,
        )
        resp.raise_for_status()
        return resp
    except requests.exceptions.RequestException:
        raise BackendError("Could not complete that request. Please try again.")


# =============================================================================
# 1) TUTOR  (RAG over course materials + guardrailed LLM)
# =============================================================================
def chat_with_tutor(user_id: str, course: str, messages: list[dict]):
    """
    Ask the tutor a question.

    Args:
        user_id:  who is asking (use it to find THEIR uploaded materials).
        course:   the course selected in the dropdown.
        messages: full chat history, a list like
                  [{"role": "user" | "assistant", "content": "..."}].
                  The last item is the newest user question.

    Returns:
        (sources, token_stream)
        sources      - list of dicts: {"file": str, "page": int, "snippet": str}
        token_stream - generator yielding pieces of the answer text.

    Why a tuple? In RAG, retrieval happens BEFORE generation, so the sources are
    known immediately and the answer streams afterwards.
    """
    _check_user(user_id)
    if not messages or messages[-1]["role"] != "user":
        raise BackendError("Please type a question for the tutor.")

    if config.USE_MOCK_BACKEND:
        question = messages[-1]["content"]

        # Pretend to retrieve: use the user's real uploaded file names if they have any.
        with _MOCK_LOCK:
            mine = list(_MOCK_MATERIALS.get(user_id, []))
        if mine:
            sources = [
                {"file": m["name"], "page": 1 + i * 3,
                 "snippet": "Relevant passage retrieved from your uploaded material."}
                for i, m in enumerate(mine[:3])
            ]
        else:
            sources = _DEFAULT_SOURCES[:2]

        answer = (
            f"Great question about **{course.split(':')[0]}**. Let's work through "
            f"\"{question[:80]}\" step by step.\n\n"
            "1. **Start with the core idea.** Try restating the problem in your own "
            "words. What are the inputs, and what output do you expect?\n"
            "2. **Connect it to the lecture.** The course materials below cover the "
            "key definition you need.\n"
            "3. **Try one small example.** Work a tiny case by hand and tell me where "
            "you get stuck - I'll give a hint rather than the full answer.\n\n"
            "What do you think the first step should be?"
        )
        return sources, _stream_text(answer)

    if config.USE_OPENAI:
        context, sources = _material_context(user_id, course)
        system = TUTOR_SYSTEM_PROMPT.format(course=course)
        if context:
            system += "\n\n## Course materials uploaded by the student\n" + context
        stream = _openai_stream([{"role": "system", "content": system}, *messages])
        return sources, stream

    # ---- REAL BACKEND -------------------------------------------------------
    # TODO (tutor agent teammate): POST to your LangGraph tutoring agent.
    #   The agent should: (1) retrieve top-k chunks for `course` from the vector
    #   store, (2) apply tutoring guardrails (guide, don't just give answers),
    #   (3) stream tokens from llama.cpp / SGLang on JetStream2.
    #   Expected response: JSON {"sources": [...]} on the first line, then
    #   streamed text. Adjust this parsing to whatever your API actually does.
    resp = _http_post(
        "/tutor/chat",
        json={"user_id": user_id, "course": course, "messages": messages},
        stream=True,
    )
    import json

    lines = resp.iter_lines(decode_unicode=True)
    first = next(lines, "{}")
    sources = json.loads(first).get("sources", [])

    def _stream() -> Iterator[str]:
        for line in lines:
            if line:
                yield line + "\n"

    return sources, _stream()


# =============================================================================
# 2) RESOURCE MANAGEMENT  (scheduling / pacing agent)
# =============================================================================
def plan_resources(user_id: str, tasks: list[dict], message: str):
    """
    Chat with the resource-management assistant, or generate a study plan.

    Args:
        user_id: who is asking.
        tasks:   the planner tasks, each like
                 {"title": str, "course": str, "due": "YYYY-MM-DD", "hours": float}
        message: what the student typed (or "Generate a study plan" from the button).

    Returns:
        token_stream - generator yielding pieces of the reply text.
    """
    _check_user(user_id)
    if not message or not message.strip():
        raise BackendError("Please enter a message first.")

    if config.USE_MOCK_BACKEND:
        if not tasks:
            reply = (
                "I don't see any tasks in your planner yet. Add a few assignments or "
                "deadlines on the left and I'll build a pacing plan around them."
            )
            return _stream_text(reply)

        # Sort by due date and build a believable plan from the real task list.
        ordered = sorted(tasks, key=lambda t: t["due"])
        total_hours = sum(t["hours"] for t in ordered)
        lines = [
            f"Here's a suggested plan for your **{len(ordered)} task(s)** "
            f"(about **{total_hours:g} hours** total):\n"
        ]
        for t in ordered:
            lines.append(
                f"- **{t['title']}** ({t['course']}) - due {t['due']}: "
                f"start at least {max(1, int(t['hours'] // 2 + 1))} day(s) early, "
                f"in blocks of about {min(t['hours'], 2):g} h."
            )
        lines.append(
            "\n**Resource tips:** book office hours for the earliest deadline this "
            "week, and reserve JetStream2 compute time before any heavy experiments."
        )
        return _stream_text("\n".join(lines))

    if config.USE_OPENAI:
        if tasks:
            task_lines = "\n".join(
                f"- {t['title']} ({t['course']}), due {t['due']}, about {t['hours']:g} h"
                for t in sorted(tasks, key=lambda t: t["due"])
            )
        else:
            task_lines = "(no tasks in the planner yet)"
        system = PLANNER_SYSTEM_PROMPT.format(today=datetime.now().strftime("%A, %Y-%m-%d"))
        user_msg = f"My planner tasks:\n{task_lines}\n\n{message}"
        return _openai_stream([
            {"role": "system", "content": system},
            {"role": "user", "content": user_msg},
        ])

    # ---- REAL BACKEND -------------------------------------------------------
    # TODO (resource agent teammate): call your LangGraph scheduling agent. It can
    #   use tools for calendar data, professor office hours, and compute quotas.
    resp = _http_post(
        "/resources/plan",
        json={"user_id": user_id, "tasks": tasks, "message": message},
        stream=True,
    )
    return (line + "\n" for line in resp.iter_lines(decode_unicode=True) if line)


# =============================================================================
# 3) PERSONALIZED FEEDBACK  (RAG + rubric-aware LLM)
# =============================================================================
def get_feedback(user_id: str, file, prompt: str) -> dict:
    """
    Generate structured feedback on a student's submission.

    Args:
        user_id: who submitted it.
        file:    a Streamlit UploadedFile (has .name, .type, .getvalue()).
        prompt:  the assignment prompt and/or rubric the student pasted in.

    Returns a dict:
        {
          "summary": str,
          "strengths": [str, ...],
          "improvements": [str, ...],
          "suggestions": [str, ...],
          "sources": [{"file": str, "page": int, "snippet": str}, ...],
        }
    """
    _check_user(user_id)
    if file is None:
        raise BackendError("Please upload a file to get feedback on.")

    if config.USE_MOCK_BACKEND:
        time.sleep(2)  # pretend the LLM is thinking
        size_kb = len(file.getvalue()) / 1024
        has_prompt = bool(prompt and prompt.strip())
        return {
            "summary": (
                f"I reviewed **{file.name}** ({size_kb:.1f} KB)"
                + (" against the assignment prompt/rubric you provided." if has_prompt
                   else ". No rubric was provided, so I used general course expectations.")
            ),
            "strengths": [
                "Clear structure: the main sections are easy to follow.",
                "Good use of course terminology in the core explanation.",
                "The scope of the work matches what the assignment asks for.",
            ],
            "improvements": [
                "Some claims need evidence or a citation to course readings.",
                "The evaluation section is thin - add measurable results.",
                "A few paragraphs repeat the same point; tighten the writing.",
            ],
            "suggestions": [
                "Add a short table comparing your approach to one alternative.",
                "Define success metrics explicitly (e.g., latency, concurrency).",
                "Re-read the rubric and map each criterion to a heading.",
            ],
            "sources": _DEFAULT_SOURCES,
        }

    if config.USE_OPENAI:
        import json

        import openai

        text = _extract_text(file)
        if not text:
            raise BackendError(f"{file.name} doesn't contain any readable text.")
        user_msg = (
            f"Assignment prompt / rubric:\n{prompt.strip() if prompt else '(none provided)'}"
            f"\n\nSubmission ({file.name}):\n{text}"
        )
        try:
            resp = _client().chat.completions.create(
                model=config.OPENAI_MODEL,
                messages=[
                    {"role": "system", "content": FEEDBACK_SYSTEM_PROMPT},
                    {"role": "user", "content": user_msg},
                ],
                response_format={"type": "json_object"},
            )
        except openai.OpenAIError as err:
            raise _openai_error(err) from err
        try:
            result = json.loads(resp.choices[0].message.content or "{}")
        except json.JSONDecodeError:
            raise BackendError("The AI model returned feedback in an unexpected format. Try again.")

        def _as_list(value) -> list[str]:
            return [str(v) for v in value] if isinstance(value, list) else []

        return {
            "summary": str(result.get("summary", "")),
            "strengths": _as_list(result.get("strengths")),
            "improvements": _as_list(result.get("improvements")),
            "suggestions": _as_list(result.get("suggestions")),
            "sources": [],
        }

    # ---- REAL BACKEND -------------------------------------------------------
    # TODO (feedback agent teammate): the agent should extract text from the file,
    #   retrieve relevant course material (RAG), and have the LLM return the JSON
    #   shape documented above. Ask the LLM for JSON output and validate it.
    resp = _http_post(
        "/feedback",
        data={"user_id": user_id, "prompt": prompt or ""},
        files={"file": (file.name, file.getvalue(), file.type or "application/octet-stream")},
    )
    return resp.json()


# =============================================================================
# 4) COURSE MATERIALS  (upload / list / delete for the RAG index)
# =============================================================================
def upload_course_material(user_id: str, file, course: str | None = None) -> dict:
    """
    Store an uploaded document and (in the real backend) index it for RAG.

    Returns the new record:
        {"id": str, "name": str, "size_kb": float, "course": str, "uploaded": str}
    """
    _check_user(user_id)
    if file is None:
        raise BackendError("No file was selected.")

    # OpenAI mode has no vector DB yet, so it reuses the in-memory store and also
    # keeps the extracted text, which the tutor pastes into its prompt.
    if config.USE_MOCK_BACKEND or config.USE_OPENAI:
        if config.USE_MOCK_BACKEND:
            time.sleep(0.5)  # pretend we're chunking + embedding
        record = {
            "id": uuid.uuid4().hex[:8],
            "name": file.name,
            "size_kb": round(len(file.getvalue()) / 1024, 1),
            "course": course or "Unassigned",
            "uploaded": datetime.now().strftime("%Y-%m-%d %H:%M"),
        }
        stored = dict(record)
        if config.USE_OPENAI:
            stored["text"] = _extract_text(file)
        with _MOCK_LOCK:
            _MOCK_MATERIALS.setdefault(user_id, []).append(stored)
        return record

    # ---- REAL BACKEND -------------------------------------------------------
    # TODO (RAG teammate): save the file, split into chunks, embed, and write to
    #   the vector store tagged with user_id (and course). Return the record.
    resp = _http_post(
        "/materials",
        data={"user_id": user_id, "course": course or ""},
        files={"file": (file.name, file.getvalue(), file.type or "application/octet-stream")},
    )
    return resp.json()


def list_course_materials(user_id: str) -> list[dict]:
    """Return all materials this user has uploaded (list of record dicts)."""
    _check_user(user_id)

    if config.USE_MOCK_BACKEND or config.USE_OPENAI:
        with _MOCK_LOCK:
            # Return metadata only; the stored text stays server-side.
            return [{k: v for k, v in m.items() if k != "text"}
                    for m in _MOCK_MATERIALS.get(user_id, [])]

    # TODO (RAG teammate): query your metadata store for this user's documents.
    return _http_get("/materials", params={"user_id": user_id}).json()


def delete_course_material(user_id: str, file_id: str) -> bool:
    """Delete one material (and its vectors, in the real backend). True on success."""
    _check_user(user_id)

    if config.USE_MOCK_BACKEND or config.USE_OPENAI:
        with _MOCK_LOCK:
            items = _MOCK_MATERIALS.get(user_id, [])
            remaining = [m for m in items if m["id"] != file_id]
            deleted = len(remaining) != len(items)
            _MOCK_MATERIALS[user_id] = remaining
        return deleted

    # TODO (RAG teammate): remove the file AND its embedded chunks from the index.
    _http_delete(f"/materials/{file_id}", params={"user_id": user_id})
    return True